"""Converts the deck JSON (1280x720 px canvas, theme tokens) into a .pptx file."""
import io
import os

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

PX = 9525  # EMU per CSS pixel at 96 dpi; 1280 px = 13.333 in
ALIGN = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}
ANCHOR = {'top': MSO_ANCHOR.TOP, 'middle': MSO_ANCHOR.MIDDLE, 'bottom': MSO_ANCHOR.BOTTOM}


def E(v):
    return Emu(int(round(float(v) * PX)))


def hexcolor(theme, v):
    if not v or v == 'none':
        return None
    v = theme.get(v, v)
    if not isinstance(v, str) or not v.startswith('#'):
        return None
    v = v[1:]
    if len(v) == 3:
        v = ''.join(c * 2 for c in v)
    return RGBColor.from_string(v[:6].upper())


def alpha(el, opacity):
    if opacity is None or opacity >= 0.999:
        return
    for clr in el.iter(qn('a:srgbClr')):
        a = etree.SubElement(clr, qn('a:alpha'))
        a.set('val', str(int(max(0.0, opacity) * 100000)))


def font_name(theme, f):
    return {'heading': theme.get('headingFont'), 'body': theme.get('bodyFont'),
            'mono': theme.get('monoFont')}.get(f) or f or 'Arial'


def add_text(slide, el, theme):
    tb = slide.shapes.add_textbox(E(el['x']), E(el['y']), E(el['w']), E(el['h']))
    tb.name = el.get('name') or 'Text'
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = ANCHOR.get(el.get('valign') or 'top', MSO_ANCHOR.TOP)
    size = float(el.get('size', 24))
    lh = float(el.get('lineHeight', 1.3))
    op = float(el.get('opacity', 1))
    color = hexcolor(theme, el.get('color', 'text'))
    name = font_name(theme, el.get('font', 'body'))
    ls = float(el.get('letterSpacing') or 0)
    for i, line in enumerate(str(el.get('text', '')).split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN.get(el.get('align', 'left'), PP_ALIGN.LEFT)
        p.line_spacing = Pt(size * 0.75 * lh)
        r = p.add_run()
        r.text = line
        f = r.font
        f.name = name
        f.size = Pt(size * 0.75)
        f.bold = int(el.get('weight', 400)) >= 600
        f.italic = bool(el.get('italic'))
        if color is not None:
            f.color.rgb = color
            alpha(r._r.get_or_add_rPr(), op)
        if ls:
            r._r.get_or_add_rPr().set('spc', str(int(round(ls * 0.75 * 100))))
    if el.get('rot'):
        tb.rotation = float(el['rot'])


def add_shape(slide, el, theme):
    kind = el.get('shape', 'rect')
    r = float(el.get('radius') or 0)
    w, h = float(el['w']), float(el['h'])
    if kind == 'ellipse':
        auto = MSO_SHAPE.OVAL
    elif kind == 'pill' or r > 0:
        auto = MSO_SHAPE.ROUNDED_RECTANGLE
    else:
        auto = MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(auto, E(el['x']), E(el['y']), E(w), E(h))
    sh.name = el.get('name') or 'Shape'
    style = sh._element.find(qn('p:style'))
    if style is not None:  # drop theme style so no default outline/shadow sneaks in
        sh._element.remove(style)
    if auto == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = 0.5 if kind == 'pill' else min(0.5, r / max(1.0, min(w, h)))
    fill = hexcolor(theme, el.get('fill'))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    stroke = hexcolor(theme, el.get('stroke'))
    sw = float(el.get('strokeWidth') or 0)
    if stroke is None or sw <= 0:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = stroke
        sh.line.width = E(sw)
    alpha(sh._element.spPr, float(el.get('opacity', 1)))
    if el.get('rot'):
        sh.rotation = float(el['rot'])


def image_path(src, uploads):
    if not src:
        return None
    p = os.path.join(uploads, os.path.basename(str(src).split('?')[0]))
    return p if os.path.isfile(p) else None


def pptx_ready(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in ('.png', '.jpg', '.jpeg', '.gif', '.bmp'):
        with Image.open(path) as im:
            return path, im.size
    if ext == '.svg':
        import cairosvg
        buf = io.BytesIO(cairosvg.svg2png(url=path, output_width=1600))
    else:
        buf = io.BytesIO()
        with Image.open(path) as im:
            im.convert('RGBA').save(buf, 'PNG')
    buf.seek(0)
    with Image.open(buf) as im:
        size = im.size
    buf.seek(0)
    return buf, size


def add_image(slide, el, uploads):
    path = image_path(el.get('src'), uploads)
    if not path:
        return
    src, (iw, ih) = pptx_ready(path)
    x, y, w, h = (float(el[k]) for k in ('x', 'y', 'w', 'h'))
    fit = el.get('fit', 'cover')
    if fit == 'contain':
        k = min(w / iw, h / ih)
        nw, nh = iw * k, ih * k
        x, y, w, h = x + (w - nw) / 2, y + (h - nh) / 2, nw, nh
    pic = slide.shapes.add_picture(src, E(x), E(y), E(w), E(h))
    pic.name = el.get('name') or 'Image'
    if fit == 'cover':
        ai, ab = iw / ih, w / h
        if ai > ab:
            c = (1 - ab / ai) / 2
            pic.crop_left = pic.crop_right = c
        elif ai < ab:
            c = (1 - ai / ab) / 2
            pic.crop_top = pic.crop_bottom = c
    op = float(el.get('opacity', 1))
    if op < 0.999:
        blip = pic._element.find('.//' + qn('a:blip'))
        if blip is not None:
            a = etree.SubElement(blip, qn('a:alphaModFix'))
            a.set('amt', str(int(op * 100000)))
    if el.get('rot'):
        pic.rotation = float(el['rot'])


def build_pptx(deck, uploads):
    theme = deck.get('theme', {})
    prs = Presentation()
    prs.slide_width = E(1280)
    prs.slide_height = E(720)
    blank = prs.slide_layouts[6]
    for s in deck.get('slides', []):
        slide = prs.slides.add_slide(blank)
        bg = hexcolor(theme, s.get('bg', 'bg'))
        if bg is not None:
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = bg
        for el in s.get('elements', []):
            if el.get('hidden'):
                continue
            t = el.get('type')
            if t == 'text':
                add_text(slide, el, theme)
            elif t == 'shape':
                add_shape(slide, el, theme)
            elif t == 'image':
                add_image(slide, el, uploads)
        if s.get('notes'):
            slide.notes_slide.notes_text_frame.text = s['notes']
    prs.core_properties.title = deck.get('title', 'Deck')
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
