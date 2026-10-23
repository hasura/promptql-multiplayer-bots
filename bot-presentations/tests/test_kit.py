"""Isolated local tests; never writes a real Google/Microsoft account."""
import copy
import importlib.util
import io
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import threading
import unittest
import urllib.error
import urllib.request
from unittest.mock import patch

APP = Path(__file__).resolve().parents[1] / 'app'
sys.path.insert(0, str(APP))
import layouts
import push
from export_pptx import build_pptx
from pptx import Presentation


class KitTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.app = Path(cls.temp.name) / 'app'
        shutil.copytree(APP, cls.app, ignore=shutil.ignore_patterns('.venv', '__pycache__', 'data', 'qa', 'deck.py'))
        subprocess.run([sys.executable, 'deck_example.py'], cwd=cls.app, check=True, capture_output=True)
        spec = importlib.util.spec_from_file_location('test_slide_server', cls.app / 'server.py')
        cls.server_module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.server_module)
        cls.server_module.get_doc()
        cls.http = cls.server_module.ThreadingHTTPServer(('127.0.0.1', 0), cls.server_module.Handler)
        cls.thread = threading.Thread(target=cls.http.serve_forever, daemon=True)
        cls.thread.start()
        cls.base = 'http://127.0.0.1:' + str(cls.http.server_port)
        cls.seed = json.loads((cls.app / 'data/seed.json').read_text())

    @classmethod
    def tearDownClass(cls):
        cls.http.shutdown()
        cls.http.server_close()
        cls.thread.join()
        cls.temp.cleanup()

    def request(self, path, method='GET', body=None, headers=None):
        h = dict(headers or {})
        if isinstance(body, dict):
            body = json.dumps(body).encode()
            h['Content-Type'] = 'application/json'
        req = urllib.request.Request(self.base + path, data=body, method=method, headers=h)
        try:
            r = urllib.request.urlopen(req, timeout=10)
        except urllib.error.HTTPError as e:
            r = e
        with r:
            raw = r.read()
            return r.status, json.loads(raw) if 'application/json' in r.headers.get('Content-Type', '') and raw else raw

    def test_sample_layouts(self):
        slides = self.seed['slides']
        self.assertEqual(len(slides), 9)
        self.assertTrue(all(s['notes'] for s in slides))
        ids = [e['id'] for s in slides for e in s['elements']]
        self.assertEqual(len(ids), len(set(ids)))
        self.assertEqual(len(layouts.PRESETS), 5)
        self.assertEqual(self.seed['theme']['name'], 'sand')
        for s in slides:
            for e in s['elements']:
                if not e.get('locked'):
                    self.assertGreaterEqual(e['x'], 0)
                    self.assertGreaterEqual(e['y'], 0)
                    self.assertLessEqual(e['x'] + e['w'], 1280)
                    self.assertLessEqual(e['y'] + e['h'], 720)

    def test_pptx_native_layers_and_notes(self):
        raw = build_pptx(self.seed, str(self.app / 'data/uploads'))
        prs = Presentation(io.BytesIO(raw))
        self.assertEqual(len(prs.slides), 9)
        self.assertEqual(prs.core_properties.title, self.seed['title'])
        for slide, source in zip(prs.slides, self.seed['slides']):
            self.assertEqual(len(slide.shapes), len(source['elements']))
            self.assertEqual(slide.notes_slide.notes_text_frame.text, source['notes'])
            actual = [sh.text for sh in slide.shapes if sh.has_text_frame and sh.text]
            expected = [e['text'] for e in source['elements'] if e['type'] == 'text']
            self.assertEqual(actual, expected)

    def test_http_conflict_and_history(self):
        status, doc = self.request('/api/deck')
        self.assertEqual(status, 200)
        changed = copy.deepcopy(doc['deck'])
        changed['title'] = 'Test revision'
        status, saved = self.request('/api/deck', 'PUT', {'deck': changed, 'baseRev': doc['rev']})
        self.assertEqual(status, 200)
        self.assertEqual(saved['rev'], doc['rev'] + 1)
        status, conflict = self.request('/api/deck', 'PUT', {'deck': doc['deck'], 'baseRev': doc['rev']})
        self.assertEqual(status, 409)
        self.assertEqual(conflict['deck']['title'], 'Test revision')
        self.assertTrue(list((self.app / 'data/history').glob('deck-*.json')))

    def test_bad_deck(self):
        self.assertEqual(self.request('/api/deck', 'PUT', {'deck': {'slides': []}})[0], 400)

    def test_ready(self):
        self.assertEqual(self.request('/readyz'), (204, b''))

    def test_theme_roundtrip(self):
        t = {**layouts.PRESETS['sand'], 'name': 'u-test', 'label': 'My theme'}
        self.assertEqual(self.request('/api/themes', 'PUT', {'themes': [t]})[0], 200)
        self.assertEqual(self.request('/api/themes')[1]['themes'][0], t)
        t['accent'] = 'bad;value'
        self.assertEqual(self.request('/api/themes', 'PUT', {'themes': [t]})[0], 400)

    def test_uploads(self):
        from PIL import Image
        b = io.BytesIO()
        Image.new('RGB', (10, 10), 'teal').save(b, format='PNG')
        status, d = self.request('/api/upload', 'POST', b.getvalue(), {'Content-Type': 'image/png'})
        self.assertEqual(status, 200)
        self.assertEqual(self.request('/' + d['src'])[1], b.getvalue())
        self.assertEqual(self.request('/api/upload', 'POST', b'', {'Content-Type': 'image/png'})[0], 413)
        self.assertEqual(self.request('/api/upload', 'POST', b'bad', {'Content-Type': 'text/plain'})[0], 415)

    def test_push_refusals(self):
        self.assertEqual(self.request('/api/push/gslides', 'POST', b'')[0], 401)
        self.assertEqual(self.request('/api/push/powerpoint', 'POST', b'')[0], 401)
        self.assertEqual(self.request('/api/push/unknown', 'POST', b'')[0], 404)

    def test_mock_google_copy(self):
        with patch.object(push, '_call', return_value=(200, {}, {'id': 'test-file'})) as call:
            out = push.to_google_slides('visitor-test', 'Deck', b'PKmock')
        args = call.call_args.args
        self.assertEqual(args[:2], ('visitor-test', '__google-drive-v2'))
        self.assertIn('uploadType=multipart', args[2])
        self.assertEqual(args[3], 'POST')
        self.assertIn(b'application/vnd.google-apps.presentation', args[4])
        self.assertIn('X-PromptQL-Description', args[5])
        self.assertEqual(out['url'], 'https://docs.google.com/presentation/d/test-file/edit')

    def test_mock_powerpoint_copy(self):
        with patch.object(push, '_call', return_value=(201, {}, {'webUrl': 'https://example.test/deck'})) as call:
            out = push.to_powerpoint('visitor-test', 'A / deck', b'PKmock')
        args = call.call_args.args
        self.assertEqual(args[:2], ('visitor-test', '__microsoft-powerpoint'))
        self.assertIn('PromptQL%20Slides/', args[2])
        self.assertEqual(args[3], 'PUT')
        self.assertIn('X-PromptQL-Description', args[5])
        self.assertEqual(out['url'], 'https://example.test/deck')

    def test_proxy_uses_visitor_token(self):
        class Response(io.BytesIO):
            status = 200
            headers = {}
        with patch.object(push, 'PLATFORM', 'https://platform.test'), patch.object(push.urllib.request, 'urlopen', return_value=Response(b'{}')) as urlopen:
            push._call('visitor-only', '__google-drive-v2', 'www.googleapis.com/test', 'GET')
        self.assertEqual(urlopen.call_args.args[0].get_header('Authorization'), 'Bearer visitor-only')

    def test_missing_account_error(self):
        err = urllib.error.HTTPError('https://platform.test', 502, 'bad gateway', {}, io.BytesIO(
            b'{"error":{"message":"Google Drive is not connected","retryable":true}}'))
        with patch.object(push, 'PLATFORM', 'https://platform.test'), patch.object(push.urllib.request, 'urlopen', side_effect=err):
            with self.assertRaises(push.PushError) as ctx:
                push._call('visitor-test', '__google-drive-v2', 'www.googleapis.com/test', 'GET')
        self.assertEqual(ctx.exception.status, 409)
        self.assertIn('not connected', str(ctx.exception))

    def test_safe_filename(self):
        self.assertEqual(push.safe_name('  A / B:*? '), 'A B')
        self.assertEqual(push.safe_name(''), 'Untitled deck')
        self.assertLessEqual(len(push.safe_name('a' * 200)), 120)


if __name__ == '__main__':
    unittest.main()